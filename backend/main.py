from __future__ import annotations

import asyncio
import io
import json
import os
import shutil
import sys
import uuid
from pathlib import Path

from dotenv import load_dotenv
from fastapi import FastAPI, File, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from loguru import logger

from config import AGENT_DEFAULTS, UPLOAD_DIR, AgentConfig, AppConfig
from orchestrator import run_pipeline
import session as session_store

load_dotenv()

# ---------------------------------------------------------------------------
# Loguru setup
# ---------------------------------------------------------------------------

# Remove default stderr handler and add structured one
logger.remove()
logger.add(
    sys.stderr,
    format="<green>{time:HH:mm:ss}</green> | <level>{level: <8}</level> | <cyan>{name}</cyan>:<cyan>{function}</cyan> — <level>{message}</level>",
    level="INFO",
)
logger.add(
    "/tmp/exphil-agent.log",
    rotation="10 MB",
    retention="7 days",
    format="{time:YYYY-MM-DD HH:mm:ss} | {level: <8} | {name}:{function}:{line} — {message}",
    level="DEBUG",
)

# ---------------------------------------------------------------------------
# App
# ---------------------------------------------------------------------------

app = FastAPI(title="Exphil Agent")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ---------------------------------------------------------------------------
# In-memory state
# ---------------------------------------------------------------------------

# Mutable copy of agent config so it can be changed at runtime
_agent_config: dict[str, dict] = {k: dict(v) for k, v in AGENT_DEFAULTS.items()}

# Track active WebSocket connections for progress streaming
_ws_connections: dict[str, WebSocket] = {}

# Store uploaded file paths per session
_session_files: dict[str, list[str]] = {}

# Store pipeline results for export
_pipeline_results: dict[str, dict] = {}

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

os.makedirs(UPLOAD_DIR, exist_ok=True)


async def _broadcast(session_id: str, event: dict) -> None:
    """Send a structured event to the WebSocket client for a session."""
    ws = _ws_connections.get(session_id)
    if ws:
        try:
            await ws.send_json(event)
        except Exception:
            pass


# ---------------------------------------------------------------------------
# REST endpoints
# ---------------------------------------------------------------------------


@app.post("/api/upload")
async def upload_files(files: list[UploadFile] = File(...)):
    """Accept PDF/txt uploads and store them in /tmp/uploads."""
    session_id = str(uuid.uuid4())
    session_dir = Path(UPLOAD_DIR) / session_id
    session_dir.mkdir(parents=True, exist_ok=True)

    saved: list[str] = []
    for f in files:
        dest = session_dir / f.filename
        with open(dest, "wb") as buf:
            shutil.copyfileobj(f.file, buf)
        saved.append(str(dest))

    _session_files[session_id] = saved
    logger.info("Upload: session={} files={}", session_id, [f.filename for f in files])
    return {"session_id": session_id, "files": [f.filename for f in files]}


@app.post("/api/start")
async def start_pipeline(body: dict):
    """Kick off the agent pipeline asynchronously."""
    session_id = body.get("session_id")
    if not session_id or session_id not in _session_files:
        return {"error": "Invalid or missing session_id"}

    task = body.get("task", "")
    if not task:
        return {"error": "Missing 'task' field (the assignment text)"}

    session_dir = str(Path(UPLOAD_DIR) / session_id)
    writer_config = body.get("writer_config", {})
    evaluator_config = body.get("evaluator_config", {})

    logger.info("Starting pipeline: session={} task='{}'", session_id, task[:80])

    async def _run():
        try:
            result = await run_pipeline(
                task=task,
                session_id=session_id,
                session_dir=session_dir,
                writer_config=writer_config,
                evaluator_config=evaluator_config,
                on_progress=lambda event: _broadcast(session_id, event),
            )
            # Store result for export
            _pipeline_results[session_id] = result
        except Exception as exc:
            logger.error("Pipeline error: session={} error={}", session_id, str(exc)[:200])
            await _broadcast(session_id, {
                "event": "error",
                "message": str(exc),
            })

    asyncio.create_task(_run())
    return {"status": "started", "session_id": session_id}


@app.get("/api/config")
async def get_config():
    """Return current agent configuration."""
    return _agent_config


@app.post("/api/config")
async def set_config(body: dict):
    """Update agent configuration (both runtime copy and AGENT_DEFAULTS for agents)."""
    for agent_name, settings in body.items():
        if agent_name in _agent_config:
            _agent_config[agent_name].update(settings)
        # Also update AGENT_DEFAULTS so agents pick up the changes
        if agent_name in AGENT_DEFAULTS:
            AGENT_DEFAULTS[agent_name].update(settings)
    return _agent_config


# ---------------------------------------------------------------------------
# Session persistence endpoints
# ---------------------------------------------------------------------------


@app.get("/api/session/{session_id}")
async def get_session(session_id: str):
    """Return the full persisted session state."""
    data = session_store.load(session_id)
    if data is None:
        return {"error": "Session not found"}
    return data


@app.get("/api/sessions")
async def list_sessions(limit: int = 20):
    """Return summaries of the most recent sessions."""
    return session_store.list_sessions(limit=limit)


# ---------------------------------------------------------------------------
# DOCX export
# ---------------------------------------------------------------------------


@app.get("/api/export/docx/{session_id}")
async def export_docx(session_id: str):
    """Generate and return a Word document from the pipeline result."""
    result = _pipeline_results.get(session_id)
    if not result or not result.get("current_draft"):
        # Fall back to session store
        stored = session_store.load(session_id)
        if stored and stored.get("final_draft"):
            result = stored
        else:
            return {"error": "No completed draft found for this session"}

    draft_data = result.get("current_draft") or result.get("final_draft")
    if not draft_data:
        return {"error": "No draft data available"}
    critique_history = result.get("critique_history") or result.get("critiques", [])
    token_usage = result.get("token_usage", [])
    total_cost = result.get("total_cost_usd", 0)

    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Title
        title_para = doc.add_heading(draft_data.get("title", "Uten tittel"), level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Metadata
        doc.add_paragraph(
            f"Antall ord: {draft_data.get('word_count', 'N/A')} | "
            f"Revisjon: {draft_data.get('revision_number', 1)} | "
            f"Totalkostnad: ${total_cost:.4f}",
            style="Subtitle",
        )

        doc.add_paragraph("")  # spacer

        # Essay body — split by markdown headings
        body = draft_data.get("body", "")
        for line in body.split("\n"):
            stripped = line.strip()
            if stripped.startswith("### "):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith("# "):
                # Skip — we already have the title
                pass
            elif stripped:
                doc.add_paragraph(stripped)

        # References
        refs = draft_data.get("references", [])
        if refs:
            doc.add_heading("Referanser", level=2)
            for r in refs:
                ref_text = f"{r.get('author', 'Ukjent')} ({r.get('year', '?')}). {r.get('title', '')}."
                if r.get("page"):
                    ref_text += f" s. {r['page']}."
                if r.get("url"):
                    ref_text += f" {r['url']}"
                doc.add_paragraph(ref_text, style="List Bullet")

        # Evaluation summary
        if critique_history:
            doc.add_page_break()
            doc.add_heading("Evalueringshistorikk", level=1)
            for i, crit in enumerate(critique_history, 1):
                doc.add_heading(f"Evaluering {i} — Score: {crit.get('score', 'N/A')}/10", level=2)
                doc.add_paragraph(crit.get("overall_assessment", ""))

                if crit.get("strengths"):
                    doc.add_heading("Styrker", level=3)
                    for s in crit["strengths"]:
                        doc.add_paragraph(s, style="List Bullet")

                if crit.get("weaknesses"):
                    doc.add_heading("Svakheter", level=3)
                    for w in crit["weaknesses"]:
                        doc.add_paragraph(w, style="List Bullet")

        # Token usage summary
        if token_usage:
            doc.add_page_break()
            doc.add_heading("Token-bruk og kostnad", level=1)
            for u in token_usage:
                doc.add_paragraph(
                    f"{u.get('agent', '?')}: "
                    f"{u.get('input_tokens', 0)} inn + {u.get('output_tokens', 0)} ut "
                    f"= ${u.get('cost_usd', 0):.4f} ({u.get('duration_s', 0):.1f}s)",
                    style="List Bullet",
                )
            doc.add_paragraph(f"\nTotal kostnad: ${total_cost:.4f}")

        # Write to buffer
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        filename = f"exphil-essay-{session_id[:8]}.docx"
        logger.info("DOCX export: session={} filename={}", session_id, filename)

        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    except ImportError:
        logger.error("python-docx not installed")
        return {"error": "python-docx is not installed. Run: pip install python-docx"}
    except Exception as e:
        logger.error("DOCX generation failed: {}", str(e)[:200])
        return {"error": f"DOCX generation failed: {e}"}


# ---------------------------------------------------------------------------
# PDF export
# ---------------------------------------------------------------------------


@app.get("/api/export/pdf/{session_id}")
async def export_pdf(session_id: str):
    """Generate and return a PDF document from the pipeline result."""
    result = _pipeline_results.get(session_id)
    if not result or not result.get("current_draft"):
        stored = session_store.load(session_id)
        if stored and stored.get("final_draft"):
            result = stored
        else:
            return {"error": "No completed draft found for this session"}

    draft_data = result.get("current_draft") or result.get("final_draft")
    if not draft_data:
        return {"error": "No draft data available"}

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, PageBreak,
        )

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=2.5 * cm,
            rightMargin=2.5 * cm,
            topMargin=2.5 * cm,
            bottomMargin=2.5 * cm,
        )

        styles = getSampleStyleSheet()

        # Custom styles
        title_style = ParagraphStyle(
            "EssayTitle",
            parent=styles["Title"],
            fontSize=16,
            alignment=TA_CENTER,
            spaceAfter=12,
        )
        body_style = ParagraphStyle(
            "EssayBody",
            parent=styles["Normal"],
            fontSize=11,
            leading=16.5,  # 1.5 line spacing
            alignment=TA_JUSTIFY,
            spaceAfter=8,
        )
        heading_style = ParagraphStyle(
            "EssayHeading",
            parent=styles["Heading2"],
            fontSize=13,
            spaceBefore=18,
            spaceAfter=8,
        )
        ref_style = ParagraphStyle(
            "EssayRef",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            leftIndent=1.5 * cm,
            firstLineIndent=-1.5 * cm,  # hanging indent
            spaceAfter=4,
        )
        header_style = ParagraphStyle(
            "PageHeader",
            parent=styles["Normal"],
            fontSize=9,
            alignment=TA_CENTER,
            textColor="#666666",
        )

        story = []

        # Header
        from datetime import date
        today = date.today().strftime("%d.%m.%Y")
        story.append(Paragraph(f"Examen Philosophicum — NTNU &nbsp;&nbsp;|&nbsp;&nbsp; {today}", header_style))
        story.append(Spacer(1, 24))

        # Title
        title = draft_data.get("title", "Uten tittel")
        story.append(Paragraph(title.replace("&", "&amp;").replace("<", "&lt;"), title_style))
        story.append(Spacer(1, 12))

        # Body — parse markdown-ish structure
        body = draft_data.get("body", "")
        for line in body.split("\n"):
            stripped = line.strip()
            if not stripped:
                story.append(Spacer(1, 6))
            elif stripped.startswith("### "):
                story.append(Paragraph(stripped[4:].replace("&", "&amp;").replace("<", "&lt;"), heading_style))
            elif stripped.startswith("## "):
                story.append(Paragraph(stripped[3:].replace("&", "&amp;").replace("<", "&lt;"), heading_style))
            elif stripped.startswith("# "):
                pass  # skip — title already rendered
            else:
                # Simple markdown bold/italic handling
                text = stripped.replace("&", "&amp;").replace("<", "&lt;")
                text = text.replace("**", "<b>", 1)
                while "**" in text:
                    text = text.replace("**", "</b>", 1)
                    if "**" in text:
                        text = text.replace("**", "<b>", 1)
                text = text.replace("*", "<i>", 1)
                while "*" in text:
                    text = text.replace("*", "</i>", 1)
                    if "*" in text:
                        text = text.replace("*", "<i>", 1)
                story.append(Paragraph(text, body_style))

        # References
        refs = draft_data.get("references", [])
        if refs:
            story.append(Spacer(1, 24))
            story.append(Paragraph("Referanser", heading_style))
            for r in refs:
                ref_text = f"{r.get('author', 'Ukjent')} ({r.get('year', '?')}). <i>{r.get('title', '')}</i>."
                if r.get("page"):
                    ref_text += f" s. {r['page']}."
                if r.get("url"):
                    ref_text += f" {r['url']}"
                story.append(Paragraph(ref_text, ref_style))

        # Page numbers
        def add_page_number(canvas, doc):
            canvas.saveState()
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor("#999999")
            canvas.drawCentredString(A4[0] / 2, 1.5 * cm, f"{doc.page}")
            canvas.restoreState()

        doc.build(story, onFirstPage=add_page_number, onLaterPages=add_page_number)
        buffer.seek(0)

        filename = f"exphil-essay-{session_id[:8]}.pdf"
        logger.info("PDF export: session={} filename={}", session_id, filename)

        return StreamingResponse(
            buffer,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    except ImportError:
        logger.error("reportlab not installed")
        return {"error": "reportlab is not installed. Run: pip install reportlab"}
    except Exception as e:
        logger.error("PDF generation failed: {}", str(e)[:200])
        return {"error": f"PDF generation failed: {e}"}


# ---------------------------------------------------------------------------
# Re-critique endpoint
# ---------------------------------------------------------------------------


@app.post("/api/critique")
async def re_critique(body: dict):
    """Run the critic agent on edited text without a full pipeline run."""
    session_id = body.get("session_id")
    edited_body = body.get("body", "")
    title = body.get("title", "")
    references = body.get("references", [])

    if not edited_body:
        return {"error": "Missing 'body' field"}

    try:
        from agents.critic import run_critic
        from agents.writer import EssayDraft, Reference
        from agents.searcher import SourceBundle

        # Build typed objects
        ref_objs = [
            Reference(
                author=r.get("author", ""),
                year=int(r.get("year", 0)),
                title=r.get("title", ""),
                page=r.get("page", ""),
                url=r.get("url"),
            )
            for r in references
        ]
        draft = EssayDraft(
            title=title,
            body=edited_body,
            references=ref_objs,
            word_count=len(edited_body.split()),
            revision_number=99,  # indicates manual re-evaluation
        )

        # Get task text from session store if available
        task_text = ""
        if session_id:
            stored = session_store.load(session_id)
            if stored:
                task_text = stored.get("task_text", "")

        # Run with empty sources (we don't need them for re-critique)
        critique_result, token_usage = await run_critic(
            task=task_text,
            draft=draft,
            sources=SourceBundle(pensum_sources=[], web_sources=[], summary="Re-evaluation of edited draft"),
        )

        critique_dict = critique_result.model_dump() if hasattr(critique_result, "model_dump") else dict(critique_result)
        logger.info("Re-critique: session={} score={}", session_id, critique_dict.get("score"))
        return {"critique": critique_dict}

    except ImportError as e:
        logger.error("Could not import critic agent: {}", str(e))
        return {"error": f"Critic agent import failed: {e}"}
    except Exception as e:
        logger.error("Re-critique failed: {}", str(e)[:200])
        return {"error": f"Re-critique failed: {e}"}


# ---------------------------------------------------------------------------
# WebSocket for progress streaming
# ---------------------------------------------------------------------------


@app.websocket("/ws/{session_id}")
async def websocket_progress(ws: WebSocket, session_id: str):
    await ws.accept()
    _ws_connections[session_id] = ws
    logger.info("WebSocket connected: session={}", session_id)
    try:
        while True:
            # Keep connection alive; client can also send messages here later
            await ws.receive_text()
    except WebSocketDisconnect:
        _ws_connections.pop(session_id, None)
        logger.info("WebSocket disconnected: session={}", session_id)
